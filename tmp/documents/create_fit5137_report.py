from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\77015\Monash\fit5137\A3\Comp\FIT5137_A3_Report.docx")
SCREENSHOT = Path(
    r"C:\Users\77015\AppData\Local\Temp\codex-clipboard-6e325e4c-705a-490e-b594-0c354b80a003.png"
)


def set_font(run, name="Arial", size=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph(paragraph, space_after=6, line_spacing=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing


if not SCREENSHOT.exists():
    raise FileNotFoundError(SCREENSHOT)
if OUTPUT.exists():
    raise FileExistsError(f"Refusing to overwrite existing report: {OUTPUT}")

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)

title_style = doc.styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(20)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

heading1 = doc.styles["Heading 1"]
heading1.font.name = "Arial"
heading1._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
heading1._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
heading1.font.size = Pt(14)
heading1.font.bold = True
heading1.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("FIT5137 Assignment 3 Report")
title.paragraph_format.space_after = Pt(18)

heading = doc.add_paragraph(style="Heading 1")
heading.add_run("Task 1 Data Restoration")
heading.paragraph_format.space_before = Pt(6)
heading.paragraph_format.space_after = Pt(8)

intro = doc.add_paragraph()
style_paragraph(intro, space_after=10)
run = intro.add_run(
    "The supplied GTFS and Australian Statistical Geography Standard datasets were "
    "restored into the ptv schema in PostgreSQL/PostGIS. Figure 1 reports the row "
    "count for each restored table using the prescribed validation query. All 11 "
    "required tables are present, including the GTFS, Mesh Block, Local Government "
    "Area, and Suburb and Locality datasets."
)
set_font(run, size=11)

figure = doc.add_paragraph()
figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
figure.paragraph_format.keep_with_next = True
figure.paragraph_format.space_before = Pt(4)
figure.paragraph_format.space_after = Pt(4)
picture = figure.add_run().add_picture(str(SCREENSHOT), width=Inches(6.9))
doc_pr = picture._inline.docPr
doc_pr.set("descr", "DBeaver query result showing row counts for all 11 tables in the ptv schema")

caption = doc.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_before = Pt(2)
caption.paragraph_format.space_after = Pt(8)
caption_run = caption.add_run(
    "Figure 1. Row counts for the 11 tables restored in the ptv schema."
)
set_font(caption_run, size=9, italic=True)

doc.core_properties.title = "FIT5137 Assignment 3 Report"
doc.core_properties.subject = "Spatial database and bus accessibility analysis"
doc.save(OUTPUT)
print(OUTPUT)
