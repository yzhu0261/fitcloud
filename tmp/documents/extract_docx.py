import sys
from docx import Document


doc = Document(sys.argv[1])
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text:
        print(text)

for table_index, table in enumerate(doc.tables, 1):
    print(f"\n[TABLE {table_index}]")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " ").strip() for cell in row.cells))

print("\n[HYPERLINK TARGETS]")
for rel in doc.part.rels.values():
    if rel.reltype.endswith("/hyperlink"):
        print(rel.target_ref)
